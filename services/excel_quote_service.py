import io
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# services/ -> sube un nivel -> static/logo_selec.png
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO_PATH = os.path.join(BASE_DIR, "static", "logo_selec.png")

COLOR_HDR = "3B2B2B"
BORDE = Border(*(Side(style="thin", color="000000"),) * 4)
CENTRO = Alignment(horizontal="center", vertical="center")
IZQ = Alignment(horizontal="left", vertical="center")
FILA_HDR = 7

COLUMNAS = [  # (letra, título, ancho) -> anchos con margen para la flecha del autofiltro
    ("B", "Fecha Actual", 18),
    ("C", "Part Number", 32),
    ("D", "Brand", 20),
    ("E", "Qty", 9),
    ("F", "Price USD", 15),
    ("G", "Delivery Time", 18),
    ("H", "Numero Registro", 21),
    ("I", "Responsable", 20),
]

COLUMNAS_TEXTO = ("C", "D")  # Part Number y Brand: texto (conserva ceros iniciales)


def _a_numero_si_corresponde(valor):
    """Convierte a int si el valor es solo dígitos; si no, lo deja como texto."""
    txt = str(valor if valor is not None else "").strip()
    return int(txt) if txt.isdigit() else txt


def generar_quote_xlsx(fecha_txt: str, filas: list) -> io.BytesIO:
    """Genera el Excel QUOTE con logo. fecha_txt en formato dd-MM-yyyy."""
    fecha = datetime.strptime(fecha_txt, "%d-%m-%Y")

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80

    ws.column_dimensions["A"].width = 3
    for letra, _, ancho in COLUMNAS:
        ws.column_dimensions[letra].width = ancho

    # Logo (alto 70 px, ancho proporcional)
    if os.path.exists(LOGO_PATH):
        logo = XLImage(LOGO_PATH)
        ratio = logo.width / logo.height
        logo.height = 70
        logo.width = int(70 * ratio)
        ws.add_image(logo, "B1")

    # Título
    ws.merge_cells("E4:H4")
    ws["E4"] = "QUOTE"
    ws["E4"].font = Font(name="Calibri", size=12, bold=True)
    ws["E4"].alignment = CENTRO

    # Encabezado
    ws.row_dimensions[FILA_HDR].height = 18
    for letra, titulo, _ in COLUMNAS:
        c = ws[f"{letra}{FILA_HDR}"]
        c.value = titulo
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLOR_HDR)
        c.alignment = CENTRO
        c.border = BORDE

    # Datos
    for i, f in enumerate(filas, start=FILA_HDR + 1):
        valores = [
            fecha,
            f.get("np", ""),
            f.get("marca", ""),
            f.get("qty", ""),
            "",  # Price USD (lo completa el proveedor)
            "",  # Delivery Time
            _a_numero_si_corresponde(f.get("registro", "")),
            f.get("responsable", ""),
        ]
        for (letra, _, _), valor in zip(COLUMNAS, valores):
            c = ws[f"{letra}{i}"]
            c.value = valor
            c.border = BORDE
            c.font = Font(name="Calibri", size=10)
            c.alignment = IZQ if letra == "C" else CENTRO
            if letra == "B":
                c.number_format = "dd-mm-yyyy"
            elif letra in COLUMNAS_TEXTO:
                c.number_format = "@"

    ultima = FILA_HDR + max(len(filas), 1)
    ws.freeze_panes = f"A{FILA_HDR + 1}"
    ws.auto_filter.ref = f"B{FILA_HDR}:I{ultima}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf